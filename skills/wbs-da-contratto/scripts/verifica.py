#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Controlli di quadratura su una WBS generata. Va eseguito su una COPIA ricalcolata
del file (le formule scritte da openpyxl non hanno valori in cache finché Excel o
LibreOffice non le valutano):

    cp file.xlsx /tmp/verifica.xlsx
    python3 .../xlsx/scripts/recalc.py /tmp/verifica.xlsx 300
    python3 verifica.py /tmp/verifica.xlsx --spec spec.json

Una ricalcolata senza errori dimostra che le formule valutano, non che i numeri
sono giusti: questi controlli servono a quello. Leggi references/controlli.md per
capire cosa significa ogni scostamento.
"""
import argparse
import json
import sys

from openpyxl import load_workbook

ap = argparse.ArgumentParser(description=__doc__)
ap.add_argument("file", help="la copia RICALCOLATA del workbook")
ap.add_argument("--spec", help="la specifica JSON, per i confronti col contratto")
args = ap.parse_args()

wb = load_workbook(args.file, data_only=True)
try:
    w, t, s = wb["WBS"], wb["Timeline"], wb["Riepilogo"]
except KeyError as e:
    sys.exit(f"foglio mancante: {e}")

spec = json.load(open(args.spec, encoding="utf-8")) if args.spec else {}
par = spec.get("parametri", {})
n_fasi = len(spec.get("fasi", [])) or (w.max_row - 3)
RT = 3 + n_fasi                     # riga del totale: dati da riga 3

esiti = []


def check(nome, ok, dettaglio):
    esiti.append((ok, nome, dettaglio))


def num(v):
    return v if isinstance(v, (int, float)) else 0


# 1 — le celle formula devono avere un valore: se sono None il file non è ricalcolato
if w[f"T{RT}"].value is None:
    sys.exit("il file non è stato ricalcolato: le formule non hanno valori in cache. "
             "Esegui recalc.py sulla copia prima di questo script.")

# 2 — quadratura del costo per tipologia di ora
tot, somma = num(s["C15"].value), num(s["C16"].value)
check("Quadratura costo giornate", abs(tot - somma) < 0.01,
      f"totale {tot:,.2f} € vs somma delle tre tipologie {somma:,.2f} €")

# 3 — sessioni in presenza vs trasferte incluse nel contratto
sess = num(s["C19"].value)
attese = par.get("trasferte_incluse")
if attese is not None:
    check("Sessioni in presenza vs trasferte incluse", sess == attese,
          f"calcolate {sess:g}, il contratto ne include {attese:g}")
else:
    check("Sessioni in presenza", True, f"{sess:g} (nessun riferimento contrattuale nella spec)")

# 4 — lead time vs durata dichiarata
eff, tgt = num(s["C45"].value), num(s["C47"].value)
check("Lead time vs target dichiarato", eff == tgt,
      f"effettivo {eff:g} settimane, dichiarato {tgt:g}")

# 5 — margine e tariffa media implicita
pieno, inv = num(s["C30"].value), num(s["C31"].value)
marg = num(s["C32"].value)
media = num(s["C34"].value)
tf, tb, tv = num(par.get("tariffa_front")), num(par.get("tariffa_back")), num(par.get("tariffa_viaggio"))
minima = min(x for x in (tf, tb, tv) if x) if any((tf, tb, tv)) else 0
soglia = par.get("margine_minimo", 0.10)
pct = marg / inv if inv else 0
check(f"Margine sopra il {soglia*100:.0f}%", inv > 0 and pct >= soglia,
      f"investimento {inv:,.2f} € − costo pieno {pieno:,.2f} € = {marg:,.2f} €"
      + (f" ({pct*100:.2f}%)" if inv else "")
      + ("" if pct >= soglia else "  ← il prezzo venduto non regge le giornate necessarie"))
if minima:
    check("Tariffa media implicita sopra la più bassa impostata", media >= minima,
          f"implicita {media:,.2f} €/h, tariffa più bassa impostata {minima:,.2f} €/h")

# 6 — ogni fase con incontri deve avere durata e persone, altrimenti front = 0 muto
buchi = []
for r in range(3, RT):
    nome = w[f"C{r}"].value
    inc, ore, pers = w[f"Q{r}"].value, w[f"R{r}"].value, w[f"S{r}"].value
    if inc and not (ore and pers):
        buchi.append(f"{nome} (incontri={inc}, ore={ore}, persone={pers})")
check("Fasi con incontri complete di durata e persone", not buchi,
      "; ".join(buchi) if buchi else "tutte complete")

# 7 — le giornate di back office sono stime: devono esserci e va detto
back = num(w[f"U{RT}"].value)
front = num(w[f"T{RT}"].value)
check("Back office valorizzato", back > 0,
      f"{back:g} giornate — sono STIME, il contratto non le quantifica "
      f"(front office calcolato: {front:g})")

# 8 — nessun foglio di servizio dimenticato
extra = [n for n in wb.sheetnames if n not in ("WBS", "Timeline", "Riepilogo", "Parametri", "Ruoli")]
check("Nessun foglio estraneo", not extra, ", ".join(extra) if extra else "solo i cinque fogli previsti")

larghezza = max(len(n) for _, n, _ in esiti)
print()
for ok, nome, dett in esiti:
    print(f"  {'OK  ' if ok else 'ATT.'}  {nome:<{larghezza}}  {dett}")
falliti = sum(1 for ok, _, _ in esiti if not ok)
print(f"\n  {len(esiti) - falliti}/{len(esiti)} controlli superati")
if falliti:
    print("  Gli scostamenti non sono necessariamente errori del file: "
          "leggi references/controlli.md prima di correggere.")
