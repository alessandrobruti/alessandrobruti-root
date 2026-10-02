from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
import copy

Y=RGBColor(0xFD,0xC3,0x00); K=RGBColor(0,0,0); W=RGBColor(255,255,255)
INK2=RGBColor(0x3A,0x3A,0x3A); RULE="DCDCDC"
F="Manrope"

prs=Presentation(); prs.slide_width=Inches(10); prs.slide_height=Inches(5.625)
s=prs.slides.add_slide(prs.slide_layouts[6])

def tb(x,y,w,h,runs,size,anchor=MSO_ANCHOR.TOP,color=K,align=PP_ALIGN.LEFT):
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=t.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,m,0)
    p=tf.paragraphs[0]; p.alignment=align
    for txt,b in runs:
        r=p.add_run(); r.text=txt; r.font.name=F; r.font.size=Pt(size); r.font.bold=b; r.font.color.rgb=color
    return t

def rect(x,y,w,h,fill):
    r=s.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb=fill; r.line.fill.background(); r.shadow.inherit=False
    st=r._element.find(qn('p:style'))
    if st is not None: r._element.remove(st)
    return r

# titolo + accento
rect(0.13,0.33,0.06,0.46,Y)
tb(0.27,0.30,6.53,0.54,[("Cosa vogliamo misurare e cosa ci serve",True)],20,MSO_ANCHOR.MIDDLE)
tb(0.27,0.86,8.6,0.24,[("Una richiesta unica, da validare insieme: per ogni KPI, il dato minimo e chi lo estrae. Consegna entro il ",False),("09/10",True),(".",False)],10,color=INK2)

rows=[
("1","Fatturato per canale","B2B distributori/rete agenti · GDO e retail · OEM · Amazon Vendor · Amazon Seller · e-shop","Righe fattura e note di credito 2023→oggi: cliente, canale, paese, articolo, importo netto","JGalileo","Amministrazione"),
("2","Fatturato per tipologia","Filtrazione · sottovuoto · health · kitchen; hardware vs consumabili","Stesso export + anagrafica articoli con famiglia","JGalileo","Amministrazione"),
("3","Margine per canale","Ricavo netto − costo prodotto − costi di canale (premi, fee, logistica)","Costo per articolo + premi e contributi ai clienti 2024-25","JGalileo / Excel","CFO"),
("4","Clienti B2B","Attivi, nuovi, persi, per canale e paese","Anagrafica clienti con tipologia e data del primo ordine","JGalileo","Commerciale"),
("5","Pipeline OEM","Opportunità, valore, stato, tempi di chiusura, vinte/perse","Elenco offerte e progetti OEM 2023→oggi (va bene un Excel)","Excel / email","Commerciale"),
("6","Vendite e costo reale Amazon","Per paese e prodotto: fee, logistica, resi, pubblicità","Business report, settlement, vendite Vendor, campagne Ads 2025","Seller/Vendor Central, Ads","Resp. Amazon"),
("7","E-shop: ordini e riacquisto","Ordini, valore medio, % di clienti che ricomprano i filtri","Export ordini 2023→oggi con id cliente","WooCommerce","Marketing"),
("8","Contatti utilizzabili","Con consenso, per fonte e paese","Export contatti + tipo di licenza","HubSpot","Marketing / IT"),
("9","Richieste clienti","Per canale e motivo, tempi di risposta","Export richieste 2025 (casella email, form o strumento in uso)","Da capire","Customer Service"),
("10","Obiettivi 2026-28","Target ricavi per canale e paese, budget marketing","Piano triennale / budget","Documento","CEO / CFO"),
]
hdr=("#","KPI da tracciare","Come lo vogliamo vedere","Dato da estrarre","Da dove","Chi")
widths=[0.30,1.55,2.55,2.95,1.20,1.05]
X0,Y0=0.20,1.22; HH,RH=0.26,0.32
gt=s.shapes.add_table(len(rows)+1,6,Inches(X0),Inches(Y0),Inches(sum(widths)),Inches(HH+RH*len(rows)))
tbl=gt.table
# niente stile di tema
tblPr=tbl._tbl.tblPr
for a in ('firstRow','bandRow'): tblPr.set(a,'0')
sid=tblPr.find(qn('a:tableStyleId'))
if sid is not None: sid.text='{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # No Style, No Grid
for i,w in enumerate(widths): tbl.columns[i].width=Inches(w)
tbl.rows[0].height=Inches(HH)
for r in range(1,len(rows)+1): tbl.rows[r].height=Inches(RH)

def border(cell,side,color,w=0.75):
    tcPr=cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn(f'a:ln{side}')): tcPr.remove(old)
    ln=etree.SubElement(tcPr,qn(f'a:ln{side}'),w=str(int(Pt(w))),cap='flat',cmpd='sng',algn='ctr')
    if color is None: etree.SubElement(ln,qn('a:noFill'))
    else:
        sf=etree.SubElement(ln,qn('a:solidFill')); etree.SubElement(sf,qn('a:srgbClr'),val=color)
    # ordine schema: lnL lnR lnT lnB prima del fill
    fill=[e for e in tcPr if e.tag in (qn('a:solidFill'),qn('a:noFill'))]
    for f_ in fill: tcPr.remove(f_); tcPr.append(f_)

def setcell(cell,txt,size,bold=False,color=K,fill=W,align=PP_ALIGN.LEFT):
    cell.fill.solid(); cell.fill.fore_color.rgb=fill
    cell.margin_left=cell.margin_right=Inches(0.06); cell.margin_top=cell.margin_bottom=Inches(0.02)
    cell.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf=cell.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=txt; r.font.name=F; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color

for c,h in enumerate(hdr):
    cell=tbl.cell(0,c); setcell(cell,h,7.5,True,W,K,PP_ALIGN.CENTER if c==0 else PP_ALIGN.LEFT)
    for sd in 'LRTB': border(cell,sd,None)
for r,row in enumerate(rows,1):
    for c,v in enumerate(row):
        cell=tbl.cell(r,c)
        if c==0: setcell(cell,v,8,True,K,Y,PP_ALIGN.CENTER)
        elif c==1: setcell(cell,v,7.6,True)
        else: setcell(cell,v,6.9,False,K if c in (4,5) else INK2)
        for sd in 'LRT': border(cell,sd,None)
        border(cell,'B',"FFFFFF" if c==0 else RULE)

# piè di slide
fy=Y0+HH+RH*len(rows)+0.12
rect(0.20,fy,9.60,0.012,K)
tb(0.20,fy+0.07,9.6,0.30,[("Formato: ",True),("CSV o Excel, va bene l'estrazione grezza.   ",False),
   ("Dati dei consumatori: ",True),("va bene l'email pseudonimizzata, purché uguale in tutti i file.   ",False),
   ("Se un dato non c'è, va detto: è già un risultato.",True)],7)

notes="""COME CONDURLA
Si parte dalla colonna KPI: è questo che volete vedere? Poi, per ciascuno: il dato c'è? Chi lo estrae?

ORDINE CONSIGLIATO
1. Fatturato per canale e per tipologia. Il 95/5 oggi è dichiarato, non misurato, e non è chiaro se Amazon Vendor stia nel B2B o nell'online.
2. Margine: è la base del business case per il payback sotto i 24 mesi chiesto da Zanuso.
3. OEM: è lì che cresce la divisione, ma non esiste una pipeline.
4. Amazon: per l'incidenza reale di fee, logistica e 595 k€ di pubblicità serve il fatturato Amazon.
5. Riacquisto filtri e contatti: servono a validare l'ipotesi dei 150 k€.
6. Richieste clienti e obiettivi.

TRE DOMANDE DA CHIUDERE AL TAVOLO
- L'e-shop vale 16 ordini al mese o "centinaia di k€"?
- L'integrazione ERP-Amazon è live o a file CSV?
- La vendita di hardware sul sito, quando?

Versione completa della richiesta: Data_Request_LAica.xlsx"""
s.notes_slide.notes_text_frame.text=notes
prs.save('Data_Request_LAica_slide.pptx')
